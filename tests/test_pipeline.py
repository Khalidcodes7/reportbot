"""Tests for ReportBot's config loading, report building, and mailer fallback."""
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from reportbot.mailer import send_report
from reportbot.pipeline import ReportConfig
from reportbot.report import build_report


SAMPLE_ROWS = [
    {"title": "Book A", "price": 10.5, "rating": "Three"},
    {"title": "Book B", "price": 20.0, "rating": "Five"},
    {"title": "Book C", "price": 5.25, "rating": "One"},
]


# ---------------------------------------------------------------- build_report

def test_build_report_creates_file_with_headers(tmp_path):
    out = tmp_path / "report.xlsx"
    build_report(SAMPLE_ROWS, str(out), title="Test Report")

    assert out.exists()
    wb = load_workbook(out)
    sheet = wb["Data"]

    # Title in row 1
    assert sheet.cell(row=1, column=1).value == "Test Report"
    # Headers in row 3
    assert sheet.cell(row=3, column=1).value == "title"
    assert sheet.cell(row=3, column=2).value == "price"
    # First data row
    assert sheet.cell(row=4, column=1).value == "Book A"
    assert sheet.cell(row=4, column=2).value == 10.5


def test_build_report_empty_rows_raises(tmp_path):
    with pytest.raises(ValueError):
        build_report([], str(tmp_path / "empty.xlsx"))


def test_build_report_with_chart(tmp_path):
    out = tmp_path / "chart.xlsx"
    build_report(SAMPLE_ROWS, str(out), chart_column="price")
    wb = load_workbook(out)
    assert "Chart" in wb.sheetnames


# ------------------------------------------------------------- ReportConfig

def test_config_from_file(tmp_path):
    cfg_file = tmp_path / "test.yaml"
    cfg_file.write_text(
        """
name: test
start_urls: https://example.com
fields:
  - name: title
    selector: "css:h1"
output:
  path: out.xlsx
  title: My Report
email:
  to: someone@example.com
""",
        encoding="utf-8",
    )

    config = ReportConfig.from_file(str(cfg_file))
    assert config.name == "test"
    assert config.start_urls == ["https://example.com"]
    assert config.output_path == "out.xlsx"
    assert config.report_title == "My Report"
    assert config.email_to == "someone@example.com"


# --------------------------------------------------------------- mailer

def test_mailer_fallback_saves_eml_without_credentials(tmp_path, monkeypatch):
    # Make sure no SendGrid env vars are set for this test
    monkeypatch.delenv("SENDGRID_API_KEY", raising=False)
    monkeypatch.delenv("REPORTBOT_FROM", raising=False)

    # Create a dummy attachment
    attachment = tmp_path / "dummy.xlsx"
    attachment.write_bytes(b"fake excel bytes")

    # Run from tmp_path so the "output/" dir goes there
    monkeypatch.chdir(tmp_path)
    result = send_report(
        to_email="someone@example.com",
        subject="Test",
        body="Hello",
        attachment_path=str(attachment),
    )
    assert "would_send.eml" in result
    assert Path("output", "would_send.eml").exists()