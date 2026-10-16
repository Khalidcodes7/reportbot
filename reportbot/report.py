"""Build formatted Excel reports from scraped data."""
from __future__ import annotations

import os
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=12)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


def _auto_width(worksheet) -> None:
    """Resize columns to fit their content."""
    for column_cells in worksheet.columns:
        length = max((len(str(cell.value)) if cell.value else 0) for cell in column_cells)
        letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[letter].width = min(length + 4, 60)


def build_report(
    rows: list[dict],
    output_path: str,
    title: str = "Report",
    chart_column: Optional[str] = None,
) -> str:
    """Write rows to a formatted .xlsx file. Returns the path written."""
    if not rows:
        raise ValueError("No rows to write - the scrape returned nothing.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    headers = list(rows[0].keys())

    # Title row
    sheet.append([title])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    # Header row (row 3, leave a blank row 2)
    sheet.append([])
    sheet.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=3, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN

    # Data rows
    for row in rows:
        sheet.append([row.get(h, "") for h in headers])

    _auto_width(sheet)
    sheet.freeze_panes = "A4"

    # Optional simple bar chart
    if chart_column and chart_column in headers:
        col_idx = headers.index(chart_column) + 1
        try:
            chart = BarChart()
            chart.title = f"{chart_column} overview"
            chart.y_axis.title = chart_column
            chart.x_axis.title = headers[0]

            data = Reference(
                sheet,
                min_col=col_idx,
                min_row=3,
                max_row=min(3 + len(rows), 3 + 20),   # cap at 20 bars so it stays readable
            )
            cats = Reference(
                sheet,
                min_col=1,
                min_row=4,
                max_row=min(3 + len(rows), 3 + 20),
            )
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 10
            chart.width = 20

            chart_sheet = workbook.create_sheet("Chart")
            chart_sheet.add_chart(chart, "B2")
        except Exception:
            # Charts are a nice-to-have; never fail the whole report over one
            pass

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    workbook.save(output_path)
    return output_path